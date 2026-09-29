"""An xlsx upload as header-keyed rows, with every cell of every column kept."""

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.importer import errors
from app.importer.detect import detect_container

MAX_BYTES = 10_000_000
MAX_EXPANDED_BYTES = 100_000_000

# Spectora headers continue into a note, e.g. "Comment Type (info, limit, defect)", so columns are matched
# on the name before the parenthesis. Matching that name exactly keeps "Default Value" apart from
# "Default Value 2 (...)".
REQUIRED_COLUMNS = {
    "section": "Section Name",
    "item": "Item Name",
    "comment_name": "Comment Name",
    "comment_text": "Comment Text",
    "comment_type": "Comment Type",
}
OPTIONAL_COLUMNS = {
    "category": "Category",
    "options": "Multiple Choice Options",
    "unit_options": "Unit Type Options",
    "recommendation": "Recommendation",
    "order": "Order",
    "answer_type": "Answer Type",
    "default_value": "Default Value",
}


@dataclass(frozen=True)
class SourceRow:
    number: int  # the spreadsheet row number the inspector sees (header row = 1)
    cells: dict[str, str]  # every column, keyed by header; "" for an empty cell


@dataclass(frozen=True)
class Sheet:
    headers: tuple[str, ...]
    rows: tuple[SourceRow, ...]
    column_for: dict[str, str]  # field -> header, for the columns the importer understands
    blank_rows: tuple[int, ...]
    sheet_names: tuple[str, ...]


def read_workbook(
    data: bytes, *, max_bytes: int = MAX_BYTES, max_expanded_bytes: int = MAX_EXPANDED_BYTES
) -> Sheet:
    if len(data) > max_bytes:
        raise errors.file_too_large(len(data), max_bytes)

    container = detect_container(data)
    if container == "legacy-xls":
        raise errors.legacy_xls()
    if container == "unknown":
        raise errors.not_a_spreadsheet()

    _check_archive(data, max_expanded_bytes)
    try:
        # data_only: never hand back formula source instead of the value a cell shows.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        sheet_names = tuple(workbook.sheetnames)
        workbook.close()
    except (
        Exception
    ) as exc:  # openpyxl raises many types for damaged files; all mean the same to the inspector
        raise errors.unreadable_workbook() from exc

    if not raw_rows:
        raise errors.not_a_spectora_export(list(REQUIRED_COLUMNS.values()), [])

    width = max(len(row) for row in raw_rows)
    headers = _header_keys(raw_rows[0], width)
    column_for = _match_columns(headers)
    missing = [name for field, name in REQUIRED_COLUMNS.items() if field not in column_for]
    if missing:
        raise errors.not_a_spectora_export(missing, [h for h in headers if not h.startswith("(column ")])

    rows: list[SourceRow] = []
    blank_rows: list[int] = []
    for offset, raw in enumerate(raw_rows[1:]):
        number = offset + 2
        values = [_cell_text(value) for value in raw] + [""] * (width - len(raw))
        if all(value == "" for value in values):
            blank_rows.append(number)
            continue
        rows.append(SourceRow(number, dict(zip(headers, values, strict=True))))

    if not rows:
        raise errors.empty_export()
    return Sheet(tuple(headers), tuple(rows), column_for, tuple(blank_rows), sheet_names)


def _check_archive(data: bytes, max_expanded_bytes: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise errors.unreadable_workbook() from exc
    if sum(member.file_size for member in members) > max_expanded_bytes:
        raise errors.expands_too_large(max_expanded_bytes)
    if "xl/workbook.xml" not in {member.filename for member in members}:
        raise errors.not_a_spreadsheet()


def _header_keys(header_row: list[object], width: int) -> list[str]:
    """Header text per column, made unique so no column's cells can overwrite another's."""
    raw = [_cell_text(value) for value in header_row] + [""] * (width - len(header_row))
    counts: dict[str, int] = {}
    for text in raw:
        counts[text] = counts.get(text, 0) + 1
    keys: list[str] = []
    seen: set[str] = set()
    for index, text in enumerate(raw):
        letter = get_column_letter(index + 1)
        if text.strip() == "":
            key = f"(column {letter})"
        elif text in seen:
            key = f"{text} (column {letter})"
        else:
            key = text
        seen.add(text)
        keys.append(key)
    return keys


def base_name(header: str) -> str:
    """The header without its parenthetical note, lowercased: "Order (w/i item)" -> "order"."""
    return re.sub(r"\s*\(.*$", "", header).strip().lower()


def _match_columns(headers: list[str]) -> dict[str, str]:
    wanted = {name.lower(): field for field, name in {**REQUIRED_COLUMNS, **OPTIONAL_COLUMNS}.items()}
    column_for: dict[str, str] = {}
    for header in headers:
        field = wanted.get(base_name(header))
        if field and field not in column_for:
            column_for[field] = header
    return column_for


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)
