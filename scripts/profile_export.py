#!/usr/bin/env python3
"""Profile a Spectora "Export to spreadsheet -> Export HTML Text" file.

Prints what the file actually contains, so importer decisions rest on
measurements rather than guesses. Run it on every new export before trusting
the importer with it.

    python scripts/profile_export.py "InterNACHI Commercial Template-2026-09-28.xls"

Requires openpyxl.
"""
import collections
import io
import re
import sys
import zipfile

import openpyxl

REQUIRED = ["section name", "item name", "comment name", "comment text", "comment type"]
COMMENT_TYPES = ("info", "limit", "defect")
Counter = collections.Counter


def text(v):
    return "" if v is None else str(v)


def sniff(data):
    """The container the bytes really are, whatever the extension says."""
    if data[:4] == b"PK\x03\x04":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "legacy-xls"
    if data.lstrip()[:1] == b"<":
        return "html-or-xml"
    return "unknown"


def column(header, prefix):
    return next((i for i, h in enumerate(header) if h.startswith(prefix)), -1)


def main(path):
    data = open(path, "rb").read()
    kind = sniff(data)
    print(f"== container ==\nfile: {path}\nbytes: {len(data)}  real format: {kind}")
    if kind != "xlsx":
        print("Not an xlsx container; this profiler stops here.")
        return 1
    ext = path.rsplit(".", 1)[-1].lower()
    if ext != "xlsx":
        print(f"NOTE: extension .{ext} does not match content (xlsx). "
              "openpyxl rejects the path by extension, so the file is read as bytes.")

    raw = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    print("cell types:", dict(Counter(re.findall(r'<c [^>]*?t="(\w+)"', raw))),
          "| self-closing cells:", len(re.findall(r"<c [^>]*/>", raw)))
    print("raw '&amp;amp;' (double-escaped):", raw.count("&amp;amp;"),
          "| single '&amp;':", len(re.findall(r"&amp;(?!amp;)", raw)))

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    header = [text(h).strip().lower() for h in rows[0]]
    body = [r for r in rows[1:] if any(text(v).strip() for v in r)]
    print(f"sheets: {wb.sheetnames}  columns: {len(header)}  data rows: {len(body)}"
          f"  blank rows skipped: {len(rows) - 1 - len(body)}")
    missing = [h for h in REQUIRED if column(header, h) < 0]
    print("required headers missing:", missing or "none")
    if missing:
        return 1

    SEC, ITEM, NAME, TXT, TYPE = (column(header, h) for h in REQUIRED)
    CAT, MC, REC, ORD, ANS = (column(header, h) for h in
                              ("category", "multiple choice", "recommendation", "order", "answer type"))
    get = lambda r, i: text(r[i]) if 0 <= i < len(r) else ""

    print("\n== column fill ==")
    for ci, h in enumerate(rows[0]):
        vals = [get(r, ci) for r in body if get(r, ci).strip()]
        distinct = Counter(vals)
        flag = "EMPTY" if not vals else ("CONSTANT" if len(distinct) == 1 else "")
        sample = [k[:40] for k, _ in distinct.most_common(3)]
        print(f"{ci:2d} {text(h)[:45]:47} filled={len(vals):4d} distinct={len(distinct):4d} {flag:8} {sample if vals else ''}")

    print("\n== hierarchy ==")
    sections, items = [], []
    for r in body:
        s, it = get(r, SEC), get(r, ITEM)
        if s not in sections:
            sections.append(s)
        if (s, it) not in items:
            items.append((s, it))
    print(f"sections: {len(sections)}  items (section+item pairs): {len(items)}"
          f"  distinct item names: {len({i for _, i in items})}")
    for s in sections:
        n = sum(1 for r in body if get(r, SEC) == s)
        print(f"  {n:4d} comments  {s!r}  ({sum(1 for x, _ in items if x == s)} items)")
    shared = {k: v for k, v in Counter(i for _, i in items).items() if v > 1}
    print("item names reused across sections:", shared or "none")
    for label, key in (("sections", lambda r: get(r, SEC)), ("items", lambda r: (get(r, SEC), get(r, ITEM)))):
        runs, prev = [], object()
        for r in body:
            if key(r) != prev:
                runs.append(key(r))
                prev = key(r)
        split = {k: v for k, v in Counter(runs).items() if v > 1}
        print(f"{label} split into non-contiguous runs:", split or "none")
    orphans = [i + 2 for i, r in enumerate(body) if not get(r, SEC).strip() or not get(r, ITEM).strip()]
    print("rows missing section or item:", orphans or "none")

    print("\n== comments ==")
    print("types:", dict(Counter(get(r, TYPE) for r in body)))
    unknown = [(i + 2, get(r, TYPE)) for i, r in enumerate(body) if get(r, TYPE) not in COMMENT_TYPES]
    print("unknown/blank types:", unknown[:10] or "none")
    for col, label in ((CAT, "category"), (ANS, "answer type"), (REC, "recommendation")):
        if col >= 0:
            print(f"{label} by type:", {t: dict(Counter(get(r, col) for r in body if get(r, TYPE) == t).most_common(6))
                                        for t in COMMENT_TYPES})
    print("empty comment text by type:", dict(Counter(get(r, TYPE) for r in body if not get(r, TXT).strip())))
    if MC >= 0:
        print("rows with multiple-choice options:", sum(1 for r in body if get(r, MC).strip()),
              "| of those with empty text:", sum(1 for r in body if get(r, MC).strip() and not get(r, TXT).strip()))
    dupes = {k: v for k, v in Counter((get(r, SEC), get(r, ITEM), get(r, NAME)) for r in body).items() if v > 1}
    print("duplicate (section, item, comment name):", dupes or "none")
    if ORD >= 0:
        groups = collections.defaultdict(list)
        for r in body:
            groups[(get(r, SEC), get(r, ITEM), get(r, TYPE))].append(r[ORD] if isinstance(r[ORD], (int, float)) else 0)
        disagree = sum(1 for g in groups.values() if g != sorted(g))
        print(f"order column: values {dict(Counter(r[ORD] for r in body))}; "
              f"groups where it disagrees with file order: {disagree} of {len(groups)}")

    print("\n== rich content (comment text) ==")
    texts = [get(r, TXT) for r in body if get(r, TXT).strip()]
    plain = [t for t in texts if not t.lstrip().startswith("<")]
    print(f"html cells: {len(texts) - len(plain)}  plain-text cells: {len(plain)}")
    print("tags:", dict(Counter(t.lower() for x in texts for t in re.findall(r"<\s*([a-zA-Z0-9]+)", x))))
    print("attributes:", dict(Counter(a.lower() for x in texts for a in re.findall(r"\s([a-zA-Z-]+)=[\"']", x))))
    print("links:", sum(len(re.findall(r"<a\s", x)) for x in texts),
          "| img/video/iframe:", sum(len(re.findall(r"<(img|video|iframe)\b", x)) for x in texts))
    empty_embeds = [x for x in texts if re.search(r'<div class="[^"]*embed[^"]*"[^>]*>\s*(&nbsp;|\xa0)?\s*</div>', x)]
    print("embed wrappers with no embed inside (content lost by the export):", len(empty_embeds))
    for col, label in ((SEC, "section"), (ITEM, "item"), (NAME, "comment name"), (TXT, "comment text"), (MC, "mc options")):
        if col >= 0:
            ents = Counter(e for r in body for e in re.findall(r"&[#a-zA-Z0-9]+;", get(r, col)))
            print(f"entities left after XML decode in {label}: {dict(ents) or 'none'}")
    print("non-breaking spaces:", sum(x.count("\xa0") for x in texts), "in", sum("\xa0" in x for x in texts), "cells")
    padded = [(i + 2, repr(get(r, c))) for i, r in enumerate(body) for c in (SEC, ITEM, NAME)
              if get(r, c) != get(r, c).strip()]
    print("names with leading/trailing whitespace:", padded[:8] or "none")
    print("max length section/item/name/text:", [max(len(get(r, c)) for r in body) for c in (SEC, ITEM, NAME, TXT)])
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
